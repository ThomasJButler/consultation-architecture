"""What `consult query` and `consult export` promise: docs/02 screen 4's
per-question dashboard and step 12's workbook, both printed rather than
paged, and neither printing an answer's text (the objective in
plans/PR-09-poc-query-export-cli.md section 1).

Driven through the fixtures taken to `ready` by the worker commands, the
same objective test_cli_worker.py drives: ingest, two `worker --once`
runs, two sign-offs, two more `worker --once` runs.
"""

from __future__ import annotations

import re
from pathlib import Path
from uuid import UUID, uuid4

import psycopg
import pytest
from openpyxl import load_workbook
from psycopg.rows import DictRow

from consult import export, query, store
from consult.cli import main
from consult.config import Settings
from consult.store import PIPELINE_ROLE
from tests.test_cli_themes import ANSWER_FRAGMENTS, ingested

pytestmark = pytest.mark.db


def test_the_query_and_export_commands(
    db: psycopg.Connection[DictRow],
    db_settings: Settings,
    capsys: pytest.CaptureFixture[str],
    tmp_path: Path,
) -> None:
    ingested(db, db_settings)
    question_ids = [
        row["id"]
        for row in db.execute(
            "SELECT id FROM question WHERE kind = 'open' ORDER BY ordinal"
        ).fetchall()
    ]
    for _ in range(2):
        assert main(["worker", "--once", "--worker", "w1"], settings=db_settings) == 0
    reviewer = str(uuid4())
    for question_id in question_ids:
        assert (
            main(
                ["sign-off", str(question_id), "--reviewer", reviewer, "--expect-version", "0"],
                settings=db_settings,
            )
            == 0
        )
    for _ in range(2):
        assert main(["worker", "--once", "--worker", "w1"], settings=db_settings) == 0
    reason_row = db.execute("SELECT id FROM question WHERE column_ref = 'o_reason'").fetchone()
    consultation_row = db.execute("SELECT id FROM consultation").fetchone()
    assert reason_row is not None
    assert consultation_row is not None
    reason_id = reason_row["id"]
    consultation_id = consultation_row["id"]
    capsys.readouterr()

    # Hand count from responses.csv, the same as test_query_db.py's own
    # against d_area=Villages: 17 of the 240 respondents give a non-blank,
    # non-duplicate o_reason answer there, and the related c_route
    # distribution among them is 8 Support, 4 Oppose, 4 Not sure. The fake
    # tags every answer with the first key of the frozen shortlist,
    # ACCESS (mapping._shortlist orders by key; fake_model.good_assignments
    # takes theme_keys[0]), so ACCESS's count is the denominator under any
    # filter that still selects some of those 17.
    code = main(
        [
            "query",
            str(reason_id),
            "--filter",
            "attr:d_area=Villages",
            "--filter",
            "theme:ACCESS",
        ],
        settings=db_settings,
    )
    out = capsys.readouterr().out
    assert code == 0
    assert "of 17 respondents who answered" in out
    assert re.search(r"^\s+ACCESS\s+Access\s+17\s+100\.0%$", out, re.M)
    assert "Support 8" in out and "Oppose 4" in out and "Not sure 4" in out
    for fragment in ANSWER_FRAGMENTS:
        assert fragment not in out

    # A malformed filter is refused by its code, before any query runs,
    # and the value that triggered it never reaches the line.
    capsys.readouterr()
    code = main(
        ["query", str(reason_id), "--filter", "other:o_safety=a-secret-value"],
        settings=db_settings,
    )
    out = capsys.readouterr().out
    assert code == 2
    assert "refused: malformed_other" in out
    assert "a-secret-value" not in out

    out_path = tmp_path / "consult.xlsx"
    capsys.readouterr()

    code = main(["export", str(consultation_id), "--out", str(out_path)], settings=db_settings)

    out = capsys.readouterr().out
    assert code == 0
    assert str(out_path) in out
    assert "240 respondents" in out and "480 open answers" in out and "4 sheets" in out
    assert re.search(r"\d+ tags", out)
    for fragment in ANSWER_FRAGMENTS:
        assert fragment not in out

    # The fixture's one formula-trigger answer reads back prefixed, and
    # every cell in the workbook is a text cell (test_export.py's own
    # proof, read back here through the command rather than the function).
    workbook = load_workbook(out_path)
    responses = workbook["Responses"]
    header = [cell.value for cell in next(responses.iter_rows(max_row=1))]
    index = {name: i for i, name in enumerate(header)}
    rows = list(responses.iter_rows(min_row=2))
    by_ref = {row[index["respondent_ref"]].value: row for row in rows}
    assert by_ref["R-0057"][index["o_safety"]].value == "'=1+1"
    for sheet in workbook.worksheets:
        for sheet_row in sheet.iter_rows():
            for cell in sheet_row:
                if cell.value not in (None, ""):
                    assert cell.data_type == "s", (sheet.title, cell.coordinate)


def test_the_query_command_reads_as_the_pipeline_role(
    db: psycopg.Connection[DictRow],
    db_settings: Settings,
    capsys: pytest.CaptureFixture[str],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    # A stand-in for the query command's first read, as
    # test_cli_themes.py's test_run_job_works_as_the_pipeline_role stands
    # in for the worker's own call, wrapped rather than replaced so the
    # rest of the command still runs for real.
    ingested(db, db_settings)
    for _ in range(2):
        assert main(["worker", "--once", "--worker", "w1"], settings=db_settings) == 0
    reviewer = str(uuid4())
    open_ids = [
        row["id"]
        for row in db.execute(
            "SELECT id FROM question WHERE kind = 'open' ORDER BY ordinal"
        ).fetchall()
    ]
    for open_id in open_ids:
        assert (
            main(
                ["sign-off", str(open_id), "--reviewer", reviewer, "--expect-version", "0"],
                settings=db_settings,
            )
            == 0
        )
    for _ in range(2):
        assert main(["worker", "--once", "--worker", "w1"], settings=db_settings) == 0
    reason_row = db.execute("SELECT id FROM question WHERE column_ref = 'o_reason'").fetchone()
    assert reason_row is not None
    reason_id = reason_row["id"]
    capsys.readouterr()

    seen: list[str] = []
    original = query.theme_table

    def recording(
        conn: psycopg.Connection[DictRow],
        question_id: UUID,
        filter: query.Filter,
        *,
        department_id: UUID,
    ) -> query.ThemeTable:
        row = conn.execute("SELECT current_user AS who").fetchone()
        assert row is not None
        seen.append(str(row["who"]))
        return original(conn, question_id, filter, department_id=department_id)

    monkeypatch.setattr(query, "theme_table", recording)

    assert main(["query", str(reason_id)], settings=db_settings) == 0
    capsys.readouterr()

    # docs/06 section 2.4's backstop is the pipeline role's missing grant
    # on the vault: the query path names no vault table (query.py's own
    # module docstring), so it should read as the role that can't reach
    # the vault even if the code above it gets that wrong, not as the
    # export role, whose vault grant this path never needs. Twice: the
    # default read, then the with=duplicates one the denominator line's
    # hidden count needs.
    assert seen == [PIPELINE_ROLE, PIPELINE_ROLE]


def test_the_query_command_holds_to_the_named_department(
    db: psycopg.Connection[DictRow],
    db_settings: Settings,
    capsys: pytest.CaptureFixture[str],
) -> None:
    # docs/06 section 2: the department is the caller's. Named with
    # --department, it holds the query to that department, so another
    # department's question reaches no respondent; left out, the command
    # takes the question's own. Ingested only: the denominator needs no
    # theme, and it's test_query_db.py's hand count of 74 o_reason rows.
    ingested(db, db_settings)
    reason_row = db.execute(
        "SELECT id, department_id FROM question WHERE column_ref = 'o_reason'"
    ).fetchone()
    assert reason_row is not None
    reason_id = reason_row["id"]
    capsys.readouterr()

    def first_line(*extra: str) -> str:
        assert main(["query", str(reason_id), *extra], settings=db_settings) == 0
        return capsys.readouterr().out.splitlines()[0]

    elsewhere = first_line("--department", str(uuid4()))
    assert elsewhere == f"question {reason_id}: of 0 respondents who answered"
    named = first_line("--department", str(reason_row["department_id"]))
    assert named == (
        f"question {reason_id}: of 74 respondents who answered "
        "(147 duplicate answers hidden; add --filter with=duplicates to count them)"
    )
    assert first_line() == named


def test_a_wrong_id_is_refused_by_name_not_a_traceback(
    db: psycopg.Connection[DictRow],
    db_settings: Settings,
    capsys: pytest.CaptureFixture[str],
    tmp_path: Path,
) -> None:
    # A question id where a consultation id belongs, or the reverse, used
    # to end in a Python traceback from export.py's or cli.py's own
    # LookupError; `themes` conflated it with "no theme set yet". All
    # three now say which id they were given and which kind it wasn't.
    ingested(db, db_settings)
    question_row = db.execute("SELECT id FROM question WHERE kind = 'open' LIMIT 1").fetchone()
    consultation_row = db.execute("SELECT id FROM consultation").fetchone()
    assert question_row is not None and consultation_row is not None
    question_id = question_row["id"]
    consultation_id = consultation_row["id"]
    capsys.readouterr()

    code = main(["query", str(consultation_id)], settings=db_settings)
    out = capsys.readouterr().out
    assert code == 1
    assert out == f"question {consultation_id}: not found\n"

    out_path = tmp_path / "wrong.xlsx"
    code = main(["export", str(question_id), "--out", str(out_path)], settings=db_settings)
    out = capsys.readouterr().out
    assert code == 1
    assert out == f"consultation {question_id}: not found\n"
    assert not out_path.exists()

    code = main(["themes", str(consultation_id)], settings=db_settings)
    out = capsys.readouterr().out
    assert code == 1
    assert out == f"question {consultation_id}: not found\n"


def test_a_filter_naming_an_unknown_column_or_question_is_refused(
    db: psycopg.Connection[DictRow],
    db_settings: Settings,
    capsys: pytest.CaptureFixture[str],
) -> None:
    ingested(db, db_settings)
    reason_row = db.execute("SELECT id FROM question WHERE column_ref = 'o_reason'").fetchone()
    assert reason_row is not None
    reason_id = reason_row["id"]
    capsys.readouterr()

    # attr: names a column this consultation doesn't have: refused by
    # code before the scope CTE ever runs, not "of 0 respondents who
    # answered" at exit 0.
    code = main(
        ["query", str(reason_id), "--filter", "attr:not_a_real_column=x"], settings=db_settings
    )
    out = capsys.readouterr().out
    assert code == 2
    assert out == "refused: unknown_column\n"

    # other: names a question the same way.
    code = main(
        ["query", str(reason_id), "--filter", "other:not_a_real_question.theme=x"],
        settings=db_settings,
    )
    out = capsys.readouterr().out
    assert code == 2
    assert out == "refused: unknown_question\n"


def test_the_query_line_says_what_it_hides(
    db: psycopg.Connection[DictRow],
    db_settings: Settings,
    capsys: pytest.CaptureFixture[str],
) -> None:
    ingested(db, db_settings)
    reason_row = db.execute("SELECT id FROM question WHERE column_ref = 'o_reason'").fetchone()
    assert reason_row is not None
    reason_id = reason_row["id"]
    capsys.readouterr()

    # Hand count test_query_db.py's own test_duplicates_are_hidden_unless_asked_for
    # pins: 74 with duplicates hidden, 221 with them shown, so the default
    # scope hides 147.
    assert main(["query", str(reason_id)], settings=db_settings) == 0
    out = capsys.readouterr().out
    assert (
        f"question {reason_id}: of 74 respondents who answered "
        "(147 duplicate answers hidden; add --filter with=duplicates to count them)"
    ) in out

    assert main(["query", str(reason_id), "--filter", "with=duplicates"], settings=db_settings) == 0
    out = capsys.readouterr().out
    assert f"question {reason_id}: of 221 respondents who answered" in out
    assert "hidden" not in out


def test_the_query_commands_hidden_count_comes_from_one_snapshot(
    db: psycopg.Connection[DictRow],
    db_settings: Settings,
    capsys: pytest.CaptureFixture[str],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """`_query` works out `hidden` from two `theme_table` calls with
    `related_distribution` between them. A READ COMMITTED transaction
    gives each of those three statements its own snapshot (Postgres
    documentation, 13.2.1), so a tag retracted between the first call and
    the third moves the count. `theme:ACCESS` makes the denominator
    depend on the tag rather than only on the duplicate flag: the fake
    tags every o_reason answer, canonical and duplicate alike, with the
    shortlist's first key (the comment on test_the_query_and_export_commands,
    above), so the hidden count under this filter is the same 147 as with
    none (checked here, not assumed: test_query_db.py's own hand count is
    for the unfiltered case).
    """
    ingested(db, db_settings)
    for _ in range(2):
        assert main(["worker", "--once", "--worker", "w1"], settings=db_settings) == 0
    reviewer = str(uuid4())
    open_ids = [
        row["id"]
        for row in db.execute(
            "SELECT id FROM question WHERE kind = 'open' ORDER BY ordinal"
        ).fetchall()
    ]
    for open_id in open_ids:
        assert (
            main(
                ["sign-off", str(open_id), "--reviewer", reviewer, "--expect-version", "0"],
                settings=db_settings,
            )
            == 0
        )
    for _ in range(2):
        assert main(["worker", "--once", "--worker", "w1"], settings=db_settings) == 0
    reason_row = db.execute("SELECT id FROM question WHERE column_ref = 'o_reason'").fetchone()
    assert reason_row is not None
    reason_id = reason_row["id"]

    version_row = db.execute(
        """
        SELECT id FROM theme_set_version
         WHERE question_id = %s AND status = 'signed_off'
         ORDER BY version_no DESC LIMIT 1
        """,
        (reason_id,),
    ).fetchone()
    assert version_row is not None
    version_id = version_row["id"]
    retracted = [
        row["answer_id"]
        for row in db.execute(
            """
            SELECT at.answer_id FROM answer_theme at
              JOIN theme t ON t.id = at.theme_id
             WHERE at.theme_set_version_id = %s AND t.key = 'ACCESS'
             LIMIT 10
            """,
            (version_id,),
        ).fetchall()
    ]
    assert len(retracted) == 10

    original = query.related_distribution

    def retract_then_call(
        conn: psycopg.Connection[DictRow],
        question_id: UUID,
        filter: query.Filter,
        *,
        department_id: UUID,
    ) -> list[tuple[str, int]]:
        # A tag batch, a retraction or a sign-off committed here, between
        # the first theme_table call and the second, is what the finding
        # reproduced: the second call sees it under READ COMMITTED and the
        # first doesn't, so the printed count moves.
        with store.connect(db_settings, autocommit=True) as second:
            second.execute(
                "UPDATE answer_theme SET retracted_at = now() WHERE answer_id = ANY(%s)",
                (retracted,),
            )
        return original(conn, question_id, filter, department_id=department_id)

    monkeypatch.setattr(query, "related_distribution", retract_then_call)
    capsys.readouterr()

    code = main(["query", str(reason_id), "--filter", "theme:ACCESS"], settings=db_settings)
    out = capsys.readouterr().out
    assert code == 0
    assert (
        f"question {reason_id}: of 74 respondents who answered "
        "(147 duplicate answers hidden; add --filter with=duplicates to count them)"
    ) in out


def test_export_refuses_its_own_error_and_names_truncated_cells(
    db: psycopg.Connection[DictRow],
    db_settings: Settings,
    capsys: pytest.CaptureFixture[str],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    ingested(db, db_settings)
    consultation_row = db.execute("SELECT id FROM consultation").fetchone()
    assert consultation_row is not None
    consultation_id = consultation_row["id"]

    # write_workbook's own refusal (a connection busy, or a cell no escape
    # could fix) has to come out as a code and exit 1, never a traceback
    # that could carry the value it failed on.
    def refuse(*_args: object, **_kwargs: object) -> export.Exported:
        raise export.ExportError(export.ExportError.CONNECTION_BUSY)

    monkeypatch.setattr(export, "write_workbook", refuse)
    out_path = tmp_path / "refused.xlsx"
    capsys.readouterr()

    code = main(["export", str(consultation_id), "--out", str(out_path)], settings=db_settings)
    out = capsys.readouterr().out
    assert code == 1
    assert out == "refused: connection_busy\n"
    assert not out_path.exists()

    # Exported carries the count write_workbook already writes into the
    # manifest, and the line names it once it isn't 0.
    def canned(*_args: object, **_kwargs: object) -> export.Exported:
        return export.Exported(respondents=1, answers=1, tags=0, sheets=2, truncated_cells=3)

    monkeypatch.setattr(export, "write_workbook", canned)
    capsys.readouterr()

    code = main(["export", str(consultation_id), "--out", str(out_path)], settings=db_settings)
    out = capsys.readouterr().out
    assert code == 0
    assert "3 cells truncated at the cap" in out

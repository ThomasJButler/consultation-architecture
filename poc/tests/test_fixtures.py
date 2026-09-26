"""What the fixture generator promises: the real file format, fictional content.

The format is the one docs/00 describes: three sheets with their headers,
options and multi-select cells joined by commas, `-` for no answer, `N/A`
for not applicable. Everything else (the topic, the questions, the
columns, the options, the answers) is made up in the generator, and a
test here holds the committed files to what the generator writes.
"""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

from make_fixture_data import (
    CLOSED_QUESTIONS,
    FIXTURES_DIR,
    FORMULA_ROW,
    OPEN_QUESTIONS,
    OUT_OF_VOCABULARY_ROWS,
    PROFORMA_REASON,
    PROFORMA_ROWS,
    SEED,
    main,
    write_fixtures,
)

SHEETS = {
    "Demographic questions": ["column_reference", "question_text"],
    "Closed questions": ["column_reference", "question_text", "response_type", "options"],
    "Open questions": ["column_reference", "question_text", "related_closed_column"],
}
RESPONSE_TYPES = {"single-select", "likert-5", "multi-select"}


def rows_of(workbook: Workbook, sheet: str) -> list[list[str]]:
    return [
        ["" if cell is None else str(cell) for cell in row]
        for row in workbook[sheet].iter_rows(min_row=2, values_only=True)
    ]


def read_csv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    with path.open(newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        assert reader.fieldnames is not None
        return list(reader.fieldnames), list(reader)


def test_the_fixture_generator_writes_the_three_sheets_and_a_responses_file(
    tmp_path: Path,
) -> None:
    written = write_fixtures(tmp_path)
    workbook = load_workbook(written.definition)

    assert workbook.sheetnames == list(SHEETS)
    for sheet, headers in SHEETS.items():
        assert [cell.value for cell in workbook[sheet][1]] == headers

    closed = rows_of(workbook, "Closed questions")
    assert {row[2] for row in closed} == RESPONSE_TYPES

    # Options are comma-joined in the workbook, and one option contains a
    # comma, so splitting the cell on commas gives one option too many. That
    # is the observation in docs/00 that puts configuration in the app.
    with_comma = [(q.column_ref, o) for q in CLOSED_QUESTIONS for o in q.options if "," in o]
    assert len(with_comma) == 1
    (column_ref, option), *_ = with_comma
    (cell,) = [row[3] for row in closed if row[0] == column_ref]
    assert option in cell
    assert (
        len(cell.split(","))
        == len(next(q for q in CLOSED_QUESTIONS if q.column_ref == column_ref).options) + 1
    )

    # One open question follows up a closed one; the other has no related
    # question, written as `-` (docs/00).
    opened = rows_of(workbook, "Open questions")
    related = [row[2] for row in opened]
    closed_refs = {row[0] for row in closed}
    assert sum(ref in closed_refs for ref in related) == 1
    assert related.count("-") == len(opened) - 1
    assert "{answer}" in next(row[1] for row in opened if row[2] in closed_refs)

    header, responses = read_csv(written.responses)
    defined = [row[0] for sheet in SHEETS for row in rows_of(workbook, sheet)]
    assert set(defined) < set(header)
    assert header[0] == "respondent_ref"
    assert len(responses) == len({row["respondent_ref"] for row in responses})

    # The two markers, and a multi-select cell carrying the comma option
    # alongside another chosen option.
    assert any(value == "-" for row in responses for value in row.values())
    assert any(row["d_commute"] == "N/A" for row in responses)
    multi = next(q for q in CLOSED_QUESTIONS if q.response_type == "multi-select")
    assert any(
        option in row[multi.column_ref] and row[multi.column_ref] != option for row in responses
    )
    assert {q.column_ref for q in OPEN_QUESTIONS} <= set(header)


def test_the_committed_fixtures_are_what_the_generator_writes(tmp_path: Path) -> None:
    written = write_fixtures(tmp_path)
    committed_responses = FIXTURES_DIR / written.responses.name
    committed_definition = FIXTURES_DIR / written.definition.name

    assert committed_responses.read_bytes() == written.responses.read_bytes()
    assert committed_definition.read_bytes() == written.definition.read_bytes()


def test_the_generator_scales_deterministically(tmp_path: Path) -> None:
    """`--scale N` is N respondents from the same seeded generator, for the
    plan benchmark's 20,000-row consultation (docs/05 section 9). The
    committed 240-row files stay what the default call writes, which the
    test above holds."""
    scale = 300
    once = write_fixtures(tmp_path / "once", respondents=scale, seed=SEED)
    again = write_fixtures(tmp_path / "again", respondents=scale, seed=SEED)
    assert once.responses.read_bytes() == again.responses.read_bytes()
    assert once.definition.read_bytes() == again.definition.read_bytes()
    reseeded = write_fixtures(tmp_path / "reseeded", respondents=scale, seed=SEED + 1)
    assert reseeded.responses.read_bytes() != once.responses.read_bytes()

    # N rows under one header line, and the default still the committed
    # file's 240 (poc/README.md).
    assert len(once.responses.read_bytes().splitlines()) == scale + 1
    _, rows = read_csv(once.responses)
    _, default_rows = read_csv(write_fixtures(tmp_path / "default").responses)
    assert len(rows) == scale
    assert len(default_rows) == 240

    # The fixed cases stay on their own rows at any scale: twelve proforma
    # answers, one leading "=", and "Unsure" only where it was put.
    def refs(row_nos: frozenset[int]) -> list[str]:
        return [f"R-{row_no:04d}" for row_no in sorted(row_nos)]

    proforma = [row["respondent_ref"] for row in rows if row["o_reason"] == PROFORMA_REASON]
    formula = [row["respondent_ref"] for row in rows if row["o_safety"].startswith("=")]
    unsure = {row["respondent_ref"] for row in rows if row["c_route"] == "Unsure"}
    assert proforma == refs(PROFORMA_ROWS)
    assert formula == refs(frozenset({FORMULA_ROW}))
    assert unsure
    assert unsure <= set(refs(OUT_OF_VOCABULARY_ROWS))

    # The flag writes what the call does, into the directory it's given.
    assert main(["--scale", str(scale), "--out", str(tmp_path / "flag")]) == 0
    assert (tmp_path / "flag" / "responses.csv").read_bytes() == once.responses.read_bytes()

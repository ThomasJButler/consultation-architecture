"""What the responses reader promises.

One row per respondent and one column per question (docs/00), read the
same way from CSV and XLSX, one row at a time, with `-` and `N/A` reaching
the validator exactly as written. The reader has no opinion about what a
cell means; that's the validator's (docs/02, section 3.2).
"""

from __future__ import annotations

import csv
import inspect
from pathlib import Path

import pytest
from openpyxl import Workbook

from consult.inputs import Caps, InputError, Refusal
from consult.responses import Responses, Row

FIXTURE = Path(__file__).resolve().parent / "fixtures" / "responses.csv"
HEADER = ["respondent_ref", "d_area", "c_modes", "o_reason"]
CELLS: list[list[object]] = [
    ["R-0001", "Villages", "Cycle, Walk", "The towpath floods."],
    ["R-0002", "-", "N/A", ""],
    [42, None, "Wheelchair, mobility scooter or similar", " padded "],
]


def write_csv(path: Path) -> Path:
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(HEADER)
        for row in CELLS:
            writer.writerow(["" if cell is None else cell for cell in row])
        writer.writerow([])
    return path


def write_xlsx(path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(HEADER)
    for row in CELLS:
        sheet.append(row)
    workbook.save(path)
    return path


EXPECTED = [
    Row(
        2,
        {
            "respondent_ref": "R-0001",
            "d_area": "Villages",
            "c_modes": "Cycle, Walk",
            "o_reason": "The towpath floods.",
        },
    ),
    Row(3, {"respondent_ref": "R-0002", "d_area": "-", "c_modes": "N/A", "o_reason": ""}),
    Row(
        4,
        {
            "respondent_ref": "42",
            "d_area": "",
            "c_modes": "Wheelchair, mobility scooter or similar",
            "o_reason": "padded",
        },
    ),
]


def test_the_responses_reader_streams_rows_and_keeps_the_markers_as_written(tmp_path: Path) -> None:
    from_csv = Responses(write_csv(tmp_path / "responses.csv"))
    from_xlsx = Responses(write_xlsx(tmp_path / "responses.xlsx"))

    assert from_csv.header == from_xlsx.header == tuple(HEADER)
    assert list(from_csv.rows()) == EXPECTED
    assert list(from_xlsx.rows()) == EXPECTED
    # A generator, so a 100,000-row file is never held whole, and it can be
    # abandoned after one row without complaint.
    rows = from_csv.rows()
    assert inspect.isgenerator(rows)
    assert next(rows) == EXPECTED[0]
    rows.close()
    # And it really is one row at a time: with a cap of one row, the first
    # row still arrives before the cap refuses the second.
    capped = Responses(write_csv(tmp_path / "capped.csv"), Caps(max_rows=1)).rows()
    assert next(capped) == EXPECTED[0]
    with pytest.raises(InputError):
        next(capped)


def test_a_short_row_is_padded_and_a_long_one_is_cut_to_the_header(tmp_path: Path) -> None:
    # What a long row may carry past the header is empty cells, as a
    # trailing comma makes; a filled one there is refused
    # (test_a_data_column_with_no_header_is_refused).
    path = tmp_path / "ragged.csv"
    path.write_text("a,b\n1\n1,2,\n", encoding="utf-8")
    assert list(Responses(path).rows()) == [
        Row(2, {"a": "1", "b": ""}),
        Row(3, {"a": "1", "b": "2"}),
    ]


def test_a_data_column_with_no_header_is_refused(tmp_path: Path) -> None:
    # Cells are keyed by header, so a data cell past the header's last
    # column has no name to go under and cutting it loses it without a
    # trace. openpyxl writes no cell for a header it wasn't given, and a
    # row is as wide as its cells once the sheet's dimension is reset, so
    # the XLSX header is two wide here and the validator's "column 3 has
    # no name" never sees the third. Refused by the reader instead, for
    # CSV and XLSX alike, with the column's position and never the cell.
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["respondent_ref", "o_reason"])
    sheet.append(["R-1", "why", "stray-cell"])
    sheet.append(["R-2", "because", "stray-cell"])
    xlsx = tmp_path / "unnamed.xlsx"
    workbook.save(xlsx)
    csv_path = tmp_path / "unnamed.csv"
    csv_path.write_text(
        "respondent_ref,o_reason\nR-1,why,stray-cell\nR-2,because,stray-cell\n", encoding="utf-8"
    )
    for path in (xlsx, csv_path):
        responses = Responses(path)
        assert responses.header == ("respondent_ref", "o_reason")
        with pytest.raises(InputError) as refused:
            list(responses.rows())
        assert (refused.value.reason, refused.value.count) == (Refusal.UNNAMED_COLUMN, 3)
        assert "stray-cell" not in str(refused.value)

    # A trailing column with a header and no data in any row is kept, each
    # row padded to it.
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["respondent_ref", "o_reason", "notes_internal"])
    sheet.append(["R-1", "why"])
    sheet.append(["R-2", "because"])
    empty = tmp_path / "empty-column.xlsx"
    workbook.save(empty)
    assert list(Responses(empty).rows()) == [
        Row(2, {"respondent_ref": "R-1", "o_reason": "why", "notes_internal": ""}),
        Row(3, {"respondent_ref": "R-2", "o_reason": "because", "notes_internal": ""}),
    ]


def test_the_fixture_reads_as_the_generator_wrote_it() -> None:
    responses = Responses(FIXTURE)
    rows = list(responses.rows())
    assert responses.header[0] == "respondent_ref"
    assert len(rows) == 240
    assert rows[0].no == 2
    assert rows[0].cells["respondent_ref"] == "R-0001"
    assert rows[-1].no == 241


def test_the_csv_field_limit_is_put_back_after_a_read(tmp_path: Path) -> None:
    # csv.field_size_limit is process-wide. A reader that lowered it and
    # left it there would refuse cells for every other reader in the
    # process, the worker's included.
    import csv

    from consult.inputs import Caps

    before = csv.field_size_limit()
    path = tmp_path / "small.csv"
    path.write_text("a\n1\n", encoding="utf-8")
    list(Responses(path, Caps(max_cell_chars=50)).rows())
    assert csv.field_size_limit() == before

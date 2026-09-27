"""What the input guards promise: a hostile file is refused before its
contents are read, with a reason and a count and never the content.

THREAT_MODEL.md row 1 is the row: an XLSX is a zip of XML, so one file can
be a zip bomb or an entity-expansion attack, and a CSV can carry a cell
the size of the file. The caps are settings (consult.config), so a
reviewer can read them and the first real upload can move them.
"""

from __future__ import annotations

import csv
import re
import shutil
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook

from consult.definition import read_definition
from consult.errors import ErrorCode
from consult.inputs import Caps, InputError, Refusal
from consult.responses import Responses, Row
from tests.test_definition import GOOD, write_workbook

FIXTURES = Path(__file__).resolve().parent / "fixtures"
SMALL = Caps(max_upload_bytes=100_000_000, max_rows=1_000, max_cell_chars=2_000, max_zip_ratio=100)


def refusal_of(action: object, *args: object) -> InputError:
    with pytest.raises(InputError) as raised:
        result = action(*args) if callable(action) else action
        # A reader that streams only refuses once it is read.
        if hasattr(result, "rows"):
            list(result.rows())
    return raised.value


def test_a_file_over_the_size_cap_is_refused_unopened() -> None:
    error = refusal_of(Responses, FIXTURES / "responses.csv", Caps(max_upload_bytes=100))
    assert error.reason is Refusal.TOO_LARGE
    assert error.count == (FIXTURES / "responses.csv").stat().st_size
    assert error.code is ErrorCode.INPUT_INVALID
    assert "R-0001" not in str(error)

    error = refusal_of(read_definition, FIXTURES / "definition.xlsx", Caps(max_upload_bytes=100))
    assert error.reason is Refusal.TOO_LARGE


def test_a_zip_that_declares_far_more_than_it_holds_is_refused(tmp_path: Path) -> None:
    # Thirty megabytes of zeros compress to a few dozen kilobytes: the
    # shape of a zip bomb, caught from the central directory alone.
    path = tmp_path / "bomb.xlsx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", b"\0" * 30_000_000)
    error = refusal_of(Responses, path, SMALL)
    assert error.reason is Refusal.ZIP_RATIO
    assert error.count > SMALL.max_zip_ratio


def test_an_xlsx_that_is_not_a_zip_is_refused(tmp_path: Path) -> None:
    path = tmp_path / "plain.xlsx"
    path.write_text("respondent_ref,o_reason\nR-1,hello\n", encoding="utf-8")
    assert refusal_of(Responses, path, SMALL).reason is Refusal.NOT_A_ZIP
    assert refusal_of(read_definition, path, SMALL).reason is Refusal.NOT_A_ZIP


def test_an_entity_expansion_in_the_workbook_xml_is_refused(tmp_path: Path) -> None:
    # The fixture workbook with one sheet's XML given an entity declaration:
    # what defusedxml forbids, and what openpyxl uses it for when installed.
    source = FIXTURES / "definition.xlsx"
    path = tmp_path / "entities.xlsx"
    with (
        zipfile.ZipFile(source) as archive,
        zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for name in archive.namelist():
            data = archive.read(name)
            if name == "xl/worksheets/sheet1.xml":
                doctype = (
                    b'<!DOCTYPE x [<!ENTITY a "aaaaaaaaaa"><!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;">]>'
                )
                # openpyxl writes the sheet without an XML declaration, so
                # the DOCTYPE goes straight in front of the root element.
                data = doctype + data
            target.writestr(name, data)
    assert refusal_of(read_definition, path, SMALL).reason is Refusal.XML_FORBIDDEN
    assert refusal_of(Responses, path, SMALL).reason is Refusal.XML_FORBIDDEN


def test_a_cell_over_the_length_cap_is_refused(tmp_path: Path) -> None:
    long_cell = "x" * (SMALL.max_cell_chars + 1)
    csv_path = tmp_path / "long.csv"
    with csv_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["respondent_ref", "o_reason"])
        writer.writerow(["R-1", long_cell])
    error = refusal_of(Responses, csv_path, SMALL)
    assert error.reason is Refusal.CELL_TOO_LONG
    assert "xxx" not in str(error)

    xlsx_path = tmp_path / "long.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["respondent_ref", "o_reason"])
    sheet.append(["R-1", long_cell])
    workbook.save(xlsx_path)
    assert refusal_of(Responses, xlsx_path, SMALL).reason is Refusal.CELL_TOO_LONG


def test_a_file_over_the_row_cap_is_refused(tmp_path: Path) -> None:
    caps = Caps(max_upload_bytes=100_000_000, max_rows=5, max_cell_chars=2_000, max_zip_ratio=100)
    error = refusal_of(Responses, FIXTURES / "responses.csv", caps)
    assert error.reason is Refusal.TOO_MANY_ROWS
    assert error.count == 5

    xlsx_path = tmp_path / "responses.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    with (FIXTURES / "responses.csv").open(newline="", encoding="utf-8") as handle:
        for row in csv.reader(handle):
            sheet.append(row)
    workbook.save(xlsx_path)
    assert refusal_of(Responses, xlsx_path, caps).reason is Refusal.TOO_MANY_ROWS


def test_the_default_caps_admit_the_fixtures(tmp_path: Path) -> None:
    copied = shutil.copy(FIXTURES / "responses.csv", tmp_path / "responses.csv")
    assert len(list(Responses(Path(copied), Caps()).rows())) == 240
    assert read_definition(FIXTURES / "definition.xlsx", Caps()).column_refs
    assert Caps().max_upload_bytes >= 200 * 1024 * 1024
    assert Caps().max_rows >= 250_000
    # A file Excel itself can save is never refused on a cell's length: its
    # own limit is 32,767 characters (Microsoft Support, "Excel
    # specifications and limits", checked 25 September 2026), and docs/01
    # records single answers of 4,000 words. The width is the staging
    # table's, not Excel's: 1,599 headers beside row_no is the most a
    # Postgres table takes (test_stage.py).
    assert Caps().max_cell_chars == 32_767
    assert Caps().max_columns == 1_599


def test_malformed_xml_a_bad_encoding_and_a_missing_file_are_refused_too(tmp_path: Path) -> None:
    # "A file that isn't what its extension says" (THREAT_MODEL.md, row 1)
    # covers more than a zip bomb: a workbook whose XML is cut short, a CSV
    # in the encoding Excel on Windows writes, and a path that isn't there.
    source = FIXTURES / "definition.xlsx"
    truncated = tmp_path / "truncated.xlsx"
    with zipfile.ZipFile(source) as archive, zipfile.ZipFile(truncated, "w") as target:
        for name in archive.namelist():
            data = archive.read(name)
            if name == "xl/worksheets/sheet1.xml":
                data = data[: len(data) // 2]
            target.writestr(name, data)
    assert refusal_of(read_definition, truncated, SMALL).reason is Refusal.UNREADABLE
    assert refusal_of(Responses, truncated, SMALL).reason is Refusal.UNREADABLE

    cp1252 = tmp_path / "windows.csv"
    cp1252.write_bytes("respondent_ref,o_reason\r\nR-1,costs £4\r\n".encode("cp1252"))
    error = refusal_of(Responses, cp1252, SMALL)
    assert error.reason is Refusal.UNREADABLE
    assert "costs" not in str(error)

    assert refusal_of(Responses, tmp_path / "absent.csv", SMALL).reason is Refusal.NOT_FOUND
    assert refusal_of(read_definition, tmp_path / "absent.xlsx", SMALL).reason is Refusal.NOT_FOUND


def test_a_row_with_more_fields_than_the_column_cap_is_refused(tmp_path: Path) -> None:
    # A million distinct header names is a small file and a very large
    # report; the cap on a row's width is the third of the CSV caps
    # THREAT_MODEL.md row 1 names, beside cells and rows.
    caps = Caps(
        max_upload_bytes=100_000_000,
        max_rows=1_000,
        max_cell_chars=2_000,
        max_zip_ratio=100,
        max_columns=3,
    )
    wide = tmp_path / "wide.csv"
    wide.write_text("a,b,c,d\n1,2,3,4\n", encoding="utf-8")
    error = refusal_of(Responses, wide, caps)
    assert error.reason is Refusal.TOO_MANY_COLUMNS
    assert error.count == 4
    assert (
        refusal_of(Responses, FIXTURES / "responses.csv", caps).reason is Refusal.TOO_MANY_COLUMNS
    )

    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["a", "b", "c"])
    sheet.append([1, 2, 3, 4])
    workbook.save(tmp_path / "wide.xlsx")
    assert refusal_of(Responses, tmp_path / "wide.xlsx", caps).reason is Refusal.TOO_MANY_COLUMNS


def test_the_definition_workbook_gets_the_cell_and_width_caps_too(tmp_path: Path) -> None:
    caps = Caps(
        max_upload_bytes=100_000_000,
        max_rows=1_000,
        max_cell_chars=50,
        max_zip_ratio=100,
        max_columns=3,
    )
    workbook = Workbook()
    del workbook["Sheet"]
    sheet = workbook.create_sheet("Demographic questions")
    sheet.append(["column_reference", "question_text"])
    sheet.append(["d_area", "x" * 51])
    for name, header in (
        ("Closed questions", ["column_reference", "question_text", "response_type", "options"]),
        ("Open questions", ["column_reference", "question_text", "related_closed_column"]),
    ):
        workbook.create_sheet(name).append(header)
    workbook.save(tmp_path / "long.xlsx")
    assert refusal_of(read_definition, tmp_path / "long.xlsx", caps).reason is Refusal.CELL_TOO_LONG

    workbook = Workbook()
    del workbook["Sheet"]
    workbook.create_sheet("Demographic questions").append(["column_reference", "question_text"])
    workbook.create_sheet("Closed questions").append(
        ["column_reference", "question_text", "response_type", "options"]
    )
    workbook.create_sheet("Open questions").append(
        ["column_reference", "question_text", "related_closed_column"]
    )
    workbook.save(tmp_path / "wide.xlsx")
    assert (
        refusal_of(read_definition, tmp_path / "wide.xlsx", caps).reason is Refusal.TOO_MANY_COLUMNS
    )


def declare_dimension(source: Path, target: Path, ref: str) -> Path:
    """A copy of a workbook whose every sheet declares `ref` as its
    dimension, whatever its cells fill. The <dimension> element is the
    file's own claim about itself, so a crafted file can claim anything."""
    with (
        zipfile.ZipFile(source) as archive,
        zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as copy,
    ):
        for name in archive.namelist():
            data = archive.read(name)
            if name.startswith("xl/worksheets/"):
                data = re.sub(rb'<dimension ref="[^"]*"', f'<dimension ref="{ref}"'.encode(), data)
            copy.writestr(name, data)
    return target


def test_a_declared_dimension_cannot_widen_a_row(tmp_path: Path) -> None:
    # openpyxl's read-only iter_rows pads every row to the sheet's declared
    # dimension, and the row cap counts only rows with something in them,
    # so this 4.9 KB file, declaring A1:XFD1048576 with two headers and one
    # cell in the last row, would be read as rows of 16,384 cells each.
    # Read by what the sheet holds, the header is its two names and the
    # far cell is one row, numbered where it sits.
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["respondent_ref", "o_reason"])
    sheet.cell(row=1_048_576, column=1, value="R-far")
    workbook.save(tmp_path / "plain.xlsx")
    path = declare_dimension(tmp_path / "plain.xlsx", tmp_path / "wide.xlsx", "A1:XFD1048576")

    # Excel's own width as the cap, so what's measured is the row's width
    # and not the cap's refusal of it.
    responses = Responses(path, Caps(max_columns=16_384))
    assert responses.header == ("respondent_ref", "o_reason")
    assert list(responses.rows()) == [
        Row(1_048_576, {"respondent_ref": "R-far", "o_reason": ""}),
    ]

    # The definition workbook is read the same way, at the default caps.
    definition = write_workbook(tmp_path / "definition.xlsx", GOOD)
    declared = declare_dimension(definition, tmp_path / "declared.xlsx", "A1:XFD1048576")
    assert read_definition(declared) == read_definition(definition)

"""Write a fictional consultation in the real file format.

Only the file format is shared with any real sample (docs/00): a definition
workbook of three sheets with the headers below, options and multi-select
cells joined by commas, `-` for no answer and `N/A` for not applicable, and a
responses file with one row per respondent and one column per question. The
topic, the questions, the column references, the options and every answer
are invented here, for the same made-up consultation on a riverside cycle
route that the wireframes in docs/02 section 12 use.

Deterministic: the seed is fixed, so `make fixtures` rewrites the same files
and a test holds the committed copies to this script. Rerun it after any
change here and commit the result. `--scale N --out DIR` writes N
respondents from the same seed somewhere else, for the plan benchmark's
20,000-row consultation (docs/05 section 9); the committed files are the
default 240.

The data carries the cases later pull requests need on purpose: an option
containing a comma (the argument for configuring in the app), a follow-up
question with a placeholder for the related closed answer, `N/A` on a
demographic column, a closed value outside its option list, a campaign
proforma repeated word for word, and one answer that starts with `=`, which
the export writer must neutralise (THREAT_MODEL.md, section 4).
"""

from __future__ import annotations

import argparse
import csv
import random
import zipfile
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.writer.excel import ExcelWriter

FIXTURES_DIR = Path(__file__).resolve().parents[1] / "tests" / "fixtures"
SEED = 19
RESPONDENTS = 240
NO_ANSWER = "-"
NOT_APPLICABLE = "N/A"
# The template token for the related closed answer in a follow-up question's
# text (docs/00). A choice of this fixture, and a setting of PR-04's importer.
PLACEHOLDER = "{answer}"


@dataclass(frozen=True)
class DemographicQuestion:
    column_ref: str
    text: str
    values: tuple[str, ...]
    not_applicable_share: float = 0.0


@dataclass(frozen=True)
class ClosedQuestion:
    column_ref: str
    text: str
    response_type: str
    options: tuple[str, ...]


@dataclass(frozen=True)
class OpenQuestion:
    column_ref: str
    text: str
    related_closed_column: str | None


DEMOGRAPHIC_QUESTIONS: tuple[DemographicQuestion, ...] = (
    DemographicQuestion(
        "d_area",
        "Which part of the district do you live in?",
        ("Town centre", "Suburbs", "Villages", "Outside the district"),
    ),
    DemographicQuestion(
        "d_commute",
        "How do you usually travel to work?",
        ("Cycle", "Walk", "Bus", "Car", "Train"),
        not_applicable_share=0.15,
    ),
    DemographicQuestion(
        "d_age",
        "Which age group are you in?",
        ("18 to 24", "25 to 44", "45 to 64", "65 and over"),
    ),
)

CLOSED_QUESTIONS: tuple[ClosedQuestion, ...] = (
    ClosedQuestion(
        "c_route",
        "Do you support the proposed riverside cycle route?",
        "single-select",
        ("Support", "Oppose", "Not sure"),
    ),
    ClosedQuestion(
        "c_modes",
        "How would you use the route?",
        "multi-select",
        ("Cycle", "Walk", "Run", "Wheelchair, mobility scooter or similar", "Push a pram"),
    ),
    ClosedQuestion(
        "c_safety",
        "How safe do you feel cycling on Riverside Road today?",
        "likert-5",
        ("Very unsafe", "Unsafe", "Neither safe nor unsafe", "Safe", "Very safe"),
    ),
)

OPEN_QUESTIONS: tuple[OpenQuestion, ...] = (
    OpenQuestion(
        "o_reason",
        f"You answered '{PLACEHOLDER}' to the previous question. Why do you feel that way?",
        "c_route",
    ),
    OpenQuestion("o_safety", "What would make the route feel safer for you?", None),
)

# Columns the workbook doesn't describe, which the configure step gives a
# role to: respondent id, identity for the vault, and ignore (docs/00).
ID_COLUMN = "respondent_ref"
IDENTITY_COLUMN = "email"
IGNORE_COLUMN = "notes_internal"

REASONS: dict[str, tuple[str, ...]] = {
    "Support": (
        "It would make the school run safer for children on bikes",
        "Riverside Road is too fast for anyone to cycle on now",
        "Fewer cars on the high street would help the shops",
        "It's cheaper than widening the road and better for the air",
        "My family would finally cycle into town instead of driving",
        "Every town along the river has one except us",
    ),
    "Oppose": (
        "We'd lose the only parking near the surgery on Mill Lane",
        "The junction by the bridge is dangerous already and this adds bikes to it",
        "The towpath is underwater every winter, so the route would be shut half the year",
        "The council has better things to spend the money on",
        "Deliveries to the shops would have nowhere to stop",
        "Nobody asked the people who actually live on Mill Lane",
    ),
    "Not sure": (
        "I'd want to see how the bridge junction is handled first",
        "It depends what happens to the parking on Mill Lane",
        "The idea is good but the towpath floods, so I'm not convinced it's usable",
        "I cycle sometimes and drive sometimes, so I can see both sides",
    ),
}
REASONS["Unsure"] = REASONS["Not sure"]

SAFETY: tuple[str, ...] = (
    "Proper lighting after dark along the towpath",
    "A separate crossing at the bridge junction, not a shared one",
    "Lower speed limits on the roads that join the route",
    "Barriers between the path and the river where it narrows",
    "Keep the path clear of parked cars and bins",
    "A better surface, the current one is full of potholes",
    "Somewhere to lock a bike at the town end",
    "Signs telling drivers the path is there",
)

# A campaign proforma: the same two answers submitted word for word by a
# group of respondents, which ingest flags at answer and respondent level
# and never deletes (docs/02, section 7, decision 9).
PROFORMA_REASON = (
    "I object to the proposed riverside cycle route. It removes parking on Mill Lane, "
    "adds traffic to the bridge junction and spends money the council does not have."
)
PROFORMA_SAFETY = "Withdraw the proposal."
PROFORMA_ROWS = frozenset({17, 41, 63, 88, 102, 119, 140, 151, 173, 199, 214, 230})
# Fourteen rows with a closed value outside the option list, for the
# validator's warning-with-a-resolution (docs/02, section 3.2).
OUT_OF_VOCABULARY_ROWS = frozenset({9, 26, 44, 58, 77, 93, 110, 128, 146, 165, 182, 203, 221, 238})
# One answer starting with `=`, for the export writer to neutralise.
FORMULA_ROW = 57


@dataclass(frozen=True)
class Written:
    definition: Path
    responses: Path


def response_columns() -> list[str]:
    return [
        ID_COLUMN,
        IDENTITY_COLUMN,
        *(q.column_ref for q in DEMOGRAPHIC_QUESTIONS),
        *(q.column_ref for q in CLOSED_QUESTIONS),
        *(q.column_ref for q in OPEN_QUESTIONS),
        IGNORE_COLUMN,
    ]


def _sentence(rng: random.Random, fragments: tuple[str, ...]) -> str:
    first, second = rng.sample(fragments, 2)
    if rng.random() < 0.55:
        return f"{first}."
    return f"{first}. {second}."


def _blank_or(rng: random.Random, share: float, value: str) -> str:
    return NO_ANSWER if rng.random() < share else value


def _respondent(rng: random.Random, row_no: int) -> dict[str, str]:
    row: dict[str, str] = {
        ID_COLUMN: f"R-{row_no:04d}",
        IDENTITY_COLUMN: f"respondent{row_no:03d}@example.org",
    }
    for question in DEMOGRAPHIC_QUESTIONS:
        if rng.random() < question.not_applicable_share:
            row[question.column_ref] = NOT_APPLICABLE
        else:
            row[question.column_ref] = _blank_or(rng, 0.04, rng.choice(question.values))

    route, modes, safety = CLOSED_QUESTIONS
    stance = rng.choices(route.options, weights=(50, 35, 15))[0]
    if row_no in OUT_OF_VOCABULARY_ROWS:
        stance = "Unsure"
    row[route.column_ref] = _blank_or(rng, 0.02, stance)
    chosen = rng.sample(modes.options, rng.choice((1, 1, 2, 2, 3)))
    row[modes.column_ref] = _blank_or(rng, 0.08, ", ".join(chosen))
    row[safety.column_ref] = _blank_or(rng, 0.04, rng.choice(safety.options))

    reason, safer = OPEN_QUESTIONS
    if row_no in PROFORMA_ROWS:
        row[route.column_ref] = "Oppose"
        row[reason.column_ref] = PROFORMA_REASON
        row[safer.column_ref] = PROFORMA_SAFETY
    else:
        if row[route.column_ref] == NO_ANSWER:
            row[reason.column_ref] = NO_ANSWER
        elif rng.random() < 0.02:
            row[reason.column_ref] = NOT_APPLICABLE
        else:
            row[reason.column_ref] = _blank_or(rng, 0.06, _sentence(rng, REASONS[stance]))
        row[safer.column_ref] = _blank_or(rng, 0.06, _sentence(rng, SAFETY))
    if row_no == FORMULA_ROW:
        row[safer.column_ref] = "=1+1"

    row[IGNORE_COLUMN] = "call back" if rng.random() < 0.05 else NO_ANSWER
    return row


def _definition_workbook() -> Workbook:
    workbook = Workbook()
    del workbook["Sheet"]
    demographic = workbook.create_sheet("Demographic questions")
    demographic.append(["column_reference", "question_text"])
    for demographic_question in DEMOGRAPHIC_QUESTIONS:
        demographic.append([demographic_question.column_ref, demographic_question.text])
    closed = workbook.create_sheet("Closed questions")
    closed.append(["column_reference", "question_text", "response_type", "options"])
    for closed_question in CLOSED_QUESTIONS:
        closed.append(
            [
                closed_question.column_ref,
                closed_question.text,
                closed_question.response_type,
                ", ".join(closed_question.options),
            ]
        )
    opened = workbook.create_sheet("Open questions")
    opened.append(["column_reference", "question_text", "related_closed_column"])
    for open_question in OPEN_QUESTIONS:
        opened.append(
            [
                open_question.column_ref,
                open_question.text,
                open_question.related_closed_column or NO_ANSWER,
            ]
        )
    # Fixed so the metadata doesn't change on every run.
    stamp = datetime(2026, 9, 19, tzinfo=UTC)
    workbook.properties.creator = "make_fixture_data.py"
    workbook.properties.created = stamp
    workbook.properties.modified = stamp
    return workbook


def _fix_zip_timestamps(path: Path) -> None:
    """Rewrite the zip with one fixed timestamp per entry.

    openpyxl stamps each entry with the wall clock, so two runs never produce
    the same bytes; with the entries re-dated the file is byte-for-byte
    reproducible and the committed copy can be held to it.
    """
    with zipfile.ZipFile(path) as source:
        entries = [(info.filename, source.read(info.filename)) for info in source.infolist()]
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as target:
        for name, data in entries:
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, data)


def write_fixtures(out_dir: Path, *, respondents: int = RESPONDENTS, seed: int = SEED) -> Written:
    out_dir.mkdir(parents=True, exist_ok=True)
    definition = out_dir / "definition.xlsx"
    responses = out_dir / "responses.csv"
    # ExcelWriter directly: Workbook.save and save_workbook both re-stamp
    # `modified` with the wall clock, and a file that changes on every run
    # can't be held to its committed copy.
    with zipfile.ZipFile(definition, "w", zipfile.ZIP_DEFLATED, allowZip64=True) as archive:
        ExcelWriter(_definition_workbook(), archive).save()
    _fix_zip_timestamps(definition)
    rng = random.Random(seed)
    columns = response_columns()
    with responses.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        for row_no in range(1, respondents + 1):
            writer.writerow(_respondent(rng, row_no))
    return Written(definition=definition, responses=responses)


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Write the fictional consultation's fixtures.")
    parser.add_argument("--scale", type=int, default=RESPONDENTS, help="respondents to write")
    parser.add_argument("--out", type=Path, default=FIXTURES_DIR, help="directory to write into")
    args = parser.parse_args(argv)
    written = write_fixtures(args.out, respondents=args.scale)
    print(f"wrote {written.definition} and {written.responses}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

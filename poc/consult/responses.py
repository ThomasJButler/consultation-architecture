"""The responses file, one row at a time.

One row per respondent and one column per question (docs/00), from CSV or
XLSX, read the same way: every cell is a string, `-` and `N/A` reach the
validator exactly as written, and nothing here decides what a cell means.
A generator, because a 100,000-row file can't be parsed inside a request
and shouldn't be held whole by a job either (docs/02, step 2).
"""

from __future__ import annotations

import csv
from collections.abc import Generator, Sequence
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


@dataclass(frozen=True)
class Row:
    """One respondent: the file's row number (the header is row 1) and the
    cells by header, padded with empty strings and cut to the header."""

    no: int
    cells: dict[str, str]


def _cell(value: object) -> str:
    return "" if value is None else str(value).strip()


def _fitted(header: Sequence[str], values: Sequence[str]) -> dict[str, str]:
    padded = [*values, *[""] * (len(header) - len(values))]
    return dict(zip(header, padded[: len(header)], strict=True))


class Responses:
    def __init__(self, path: Path) -> None:
        self.path = path
        self.header = self._read_header()

    @property
    def is_xlsx(self) -> bool:
        return self.path.suffix.lower() == ".xlsx"

    def _raw_rows(self) -> Generator[tuple[str, ...], None, None]:
        if self.is_xlsx:
            workbook = load_workbook(self.path, read_only=True, data_only=True)
            try:
                for sheet_values in workbook.worksheets[0].iter_rows(values_only=True):
                    yield tuple(_cell(value) for value in sheet_values)
            finally:
                workbook.close()
        else:
            # utf-8-sig, so a byte-order mark from a spreadsheet's CSV export
            # doesn't end up glued to the first header.
            with self.path.open(newline="", encoding="utf-8-sig") as handle:
                for csv_values in csv.reader(handle):
                    yield tuple(_cell(value) for value in csv_values)

    def _read_header(self) -> tuple[str, ...]:
        rows = self._raw_rows()
        try:
            return next(rows, ())
        finally:
            rows.close()

    def rows(self) -> Generator[Row, None, None]:
        raw = self._raw_rows()
        next(raw, None)
        for offset, values in enumerate(raw, start=2):
            if not any(values):
                continue
            yield Row(offset, _fitted(self.header, values))
